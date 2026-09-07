"""Baran'ın Excel form (görsel sözleşme) — main mimarisi ile temiz port.

Bu test paketi, ``backend/server.py::export_excel`` fonksiyonunun
aşağıdaki sözleşmeyi koruduğunu kanıtlar:

* **Görsel sözleşme (Baran formuna sadık)**
  - Sheet adı ``İSG Risk Analizi (5x5)``
  - A1 başlık: ``ABCD Tech Solutions SANAYİ VE TİCARET A.Ş.`` — Calibri 16 bold #8B0000
  - A2 durum alt başlığı — Calibri 11 italic #475569
  - A4/D4/A5/D5 metadata grid (Restoran/Adres/Tarih/Denetçi)
  - Header background #8B0000, font Calibri 11 bold beyaz
  - 11 kolon sırası (No, Kategori, Tehlike & Risk Maddesi, Cevap,
    Olasılık (O), Şiddet (Ş), Risk Skoru (R), Risk Seviyesi,
    Sorumlu, Termin Süresi, Alınması Gereken Tedbir (DÖF))
  - Kolon genişlikleri ``[6, 24, 55, 12, 12, 12, 14, 18, 20, 15, 50]``
  - Filename ``ISG_Risk_Analizi_<restaurant_name>.xlsx``
  - Content-Disposition dual UTF-8

* **Veri kaynağı (mevcut main mimarisi)**
  - ``_get_audit_questions(doc)`` → audit kendi snapshot'ından beslenir
  - ``_get_audit_overrides(doc)`` → ``risk_overrides`` map
  - ``compute_effective_risk(q, override)`` → kanonik skor/level
  - ``classify_risk`` kanonik eşikleri (1-4 / 5-12 / 13-25)
  - Yeni şema primary; legacy fallback (``soru``/``sorumlu``/``termin``
    /``tedbir``/``kategori``) export'ta korunur

* **Reject edilen Baran regresyonları**
  - Türkçe → ASCII transliteration ``fix_tr_text`` YOK
  - Global ``QUESTIONS`` kullanımı YOK
  - Skor/level ``o*s`` veya ``q.get("risk_seviyesi")`` ile değil
  - Status metni ``TAMAMLANDI (RESMİ ONAYLI)`` YOK, yalnız
    ``is_completed`` bayrağına göre ``TAMAMLANDI`` / ``TASLAK DENETİM``
  - Cevap mapping: EVET→"Uygun", HAYIR→"Uygun Değil",
    NA→"N/A", empty→"Yanıtlanmadı"

Bu testler ``export_excel`` endpoint'ini end-to-end çağırır; ``StreamingResponse``
body'sini ``openpyxl`` ile yeniden açarak veriyi görsel sözleşmeye göre
doğrular. ``export_pdf``'in etkilenmediğini ayrıca bir PDF MI/scenario testiyle
kanıtlar.
"""
from __future__ import annotations

import sys
import io
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import pytest
from bson import ObjectId

from _phase_2a_helpers import (
    fake_db,
    push_audit,
    run_async,
    sample_audit_doc,
    sample_questions,
    sample_template_doc,
)


# Deterministik ObjectId — server ``ObjectId(audit_id)`` parse eder.
AID = str(ObjectId("507f1f77bcf86cd799439041"))


# ---------------------------------------------------------------------------
# Türkçe karakter yoğunluğu yüksek test fixture'ı — Baran formunun
# 'tehlike & risk maddesi' / 'sorumlu' / 'termin' / 'tedbir' kolonlarına
# Unicode karakterlerin nasıl yansıdığını gözlemlemek için.
TR_HOTSPOT = {
    "id": 1, "no": 1, "category": "Elektrik Güvenliği",
    "area": "Elektrik & Pano",
    "question": "Şişli şubesinde elektrik panosu etiketleri güncel mi?",
    "responsible": "İSG Uzmanı — Çağatay Gümüş",
    "default_probability": 5, "default_severity": 5,
    "default_risk_score": 25, "default_risk_level": "Kabul Edilemez",
    "document_risk_level": "Kabul Edilemez",
    "deadline": "2026-09-30", "legal_basis": ["İş Sağlığı ve Güvenliği Kanunu md.10"],
    "corrective_action": "Etiketleri güncelle; topraklama ölçümü yap.",
}

# q2: Dikkate Değer (3×4=12) — override farklı senaryosu için
TR_DEFAULT = {
    "id": 2, "no": 2, "category": "Yangın Güvenliği",
    "area": "Yangın Söndürme",
    "question": "Tüp doluluk oranı kontrol ediliyor mu?",
    "responsible": "İşveren Vekili",
    "default_probability": 3, "default_severity": 4,
    "default_risk_score": 12, "default_risk_level": "Dikkate Değer",
    "document_risk_level": "Dikkate Değer",
    "deadline": "2026-10-15",
    "corrective_action": "Aylık doluluk kontrolü planla.",
}

# q3: Kabul Edilebilir (1×1=1) — EVET yolu için
TR_LOW = {
    "id": 3, "no": 3, "category": "Acil Çıkış",
    "area": "Acil Çıkış & Yönlendirme",
    "question": "Acil çıkış levhaları görünür durumda mı?",
    "responsible": "Bölüm Sorumlusu",
    "default_probability": 1, "default_severity": 1,
    "default_risk_score": 1, "default_risk_level": "Kabul Edilebilir",
    "document_risk_level": "Kabul Edilebilir",
    "deadline": "2026-12-31",
    "corrective_action": "Yıllık kontrol planı.",
}

# q4: NA ile cevaplanacak satır
TR_NA = {
    "id": 4, "no": 4, "category": "Tehlikeli Kimyasal",
    "area": "Kimyasal Depolama",
    "question": "Asbest kullanımı var mı?",
    "responsible": "İSG Uzmanı",
    "default_probability": 1, "default_severity": 1,
    "default_risk_score": 1, "default_risk_level": "Kabul Edilebilir",
    "document_risk_level": "Kabul Edilebilir",
    "deadline": "",
    "corrective_action": "",
}

# q5: cevaplanmamış (missing) satır
TR_EMPTY = {
    "id": 5, "no": 5, "category": "Ergonomi",
    "area": "Ofis & Kasa",
    "question": "Sandalye ergonomik mi?",
    "responsible": "İK",
    "default_probability": 1, "default_severity": 1,
    "default_risk_score": 1, "default_risk_level": "Kabul Edilebilir",
    "document_risk_level": "Kabul Edilebilir",
    "deadline": "",
    "corrective_action": "",
}


def _tpl_tr():
    return sample_template_doc(questions=[TR_HOTSPOT, TR_DEFAULT, TR_LOW, TR_NA, TR_EMPTY])


def _tpl_legacy():
    """Legacy snapshot — ``soru``/``sorumlu``/``termin``/``tedbir``/``kategori`` alanları."""
    qs = [
        {
            "id": 1, "no": 1,
            "kategori": "Elektrik",
            "sorumlu": "İSG Uzmanı",
            "soru": "Pano topraklaması ölçüldü mü?",
            "olasilik": 3, "siddet": 4,
            "risk_skoru": 12, "risk_seviyesi": "Dikkate Değer",
            "termin": "2026-11-30",
            "tedbir": "Topraklama ölçümü planla.",
        },
    ]
    return sample_template_doc(questions=qs)


def _user():
    return {"id": "u-1", "name": "Denetçi 1"}


def _audit(*, answers=None, overrides=None, with_snapshot=True, template=None,
           is_completed=False, legacy=False):
    tpl = template or _tpl_tr()
    doc = sample_audit_doc(
        audit_id=AID,
        answers=answers or {},
        risk_overrides=overrides if overrides is not None else {},
        with_snapshot=with_snapshot,
        template=tpl,
        version=0,
    )
    doc["is_completed"] = is_completed
    # S18 lifecycle: nihai rapor export'u yalnızca DOF_CLOSED/FINAL'de
    # mümkün (DRAFT/DOF_OPEN 409). Design testleri export çıktısını
    # (layout/kolon/stil/içerik) incelediği için geçerli bir exportable
    # state sağlanır. ``is_completed`` (subtitle etiketi) ayrıca testin
    # kontrolünde kalır — state guard'ı atlanmaz, sadece geçerli state verilir.
    doc["state"] = "FINAL" if is_completed else "DOF_CLOSED"
    if legacy:
        # Legacy snapshot'ta yeni şema alanları yok; sadece eski alanlar var.
        doc["template_snapshot"] = {
            "template_code": "legacy_v0",
            "template_name": "Legacy v0",
            "template_version": 0,
            "snapshot_at": "2025-01-01T00:00:00+00:00",
            "questions": _tpl_legacy()["questions"],
        }
    return doc


def _body_bytes(response) -> bytes:
    """StreamingResponse.body → bytes.

    Starlette/FastAPI normal kullanımda .body() coroutine döner; sync test
    köprüsünde stream'i doğrudan okuyoruz.
    """
    # StreamingResponse iter chunks veya body attribute.
    if hasattr(response, "body"):
        body = response.body
        if isinstance(body, bytes):
            return body
        # Async body_iterator — sync test'te döngü kur
        if hasattr(body, "__aiter__"):
            import asyncio
            async def _collect():
                out = bytearray()
                async for chunk in body:
                    out.extend(chunk)
                return bytes(out)
            return asyncio.run(_collect())
    # Background: iter Starlette stream
    if hasattr(response, "body_iterator"):
        import asyncio
        async def _collect():
            out = bytearray()
            async for chunk in response.body_iterator:
                out.extend(chunk)
            return bytes(out)
        return asyncio.run(_collect())
    raise RuntimeError("StreamingResponse body alınamadı")


def _load_workbook(response):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(_body_bytes(response)), data_only=True)


# ===========================================================================
# 1) Görsel sözleşme — başlık / metadata / sheet adı
# ===========================================================================


class TestVisualContract:
    """Baran'ın formunun korunduğunu kanıtlar."""

    def test_sheet_name_is_in_5x5(self):
        from server import export_excel

        audit = _audit()
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        assert wb.active.title == "İSG Risk Analizi (5x5)"

    def test_a1_title_text_and_font(self):
        from server import export_excel

        audit = _audit()
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        a1 = wb.active["A1"]
        assert a1.value == "ABCD Tech Solutions SANAYİ VE TİCARET A.Ş."
        assert a1.font.name == "Calibri"
        assert a1.font.bold is True
        assert a1.font.size == 16
        # color nesnesi openpyxl'da ``rgb`` attribute'u taşır (veya ``value``)
        rgb = getattr(a1.font.color, "rgb", None) or getattr(a1.font.color, "value", None)
        assert rgb is not None
        assert str(rgb).upper().endswith("8B0000")

    def test_a2_status_subtitle_completed(self):
        from server import export_excel

        audit = _audit(is_completed=True)
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        a2 = wb.active["A2"]
        assert a2.value == "İş Sağlığı ve Güvenliği Risk Değerlendirmesi Tablosu — [TAMAMLANDI]"
        assert a2.font.name == "Calibri"
        assert a2.font.italic is True
        assert a2.font.size == 11
        rgb = getattr(a2.font.color, "rgb", None) or getattr(a2.font.color, "value", None)
        assert str(rgb).upper().endswith("475569")

    def test_a2_status_subtitle_draft(self):
        from server import export_excel

        audit = _audit(is_completed=False)
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        a2 = wb.active["A2"]
        assert a2.value == "İş Sağlığı ve Güvenliği Risk Değerlendirmesi Tablosu — [TASLAK DENETİM]"

    def test_metadata_grid_a4_d4_a5_d5(self):
        from server import export_excel

        audit = _audit()
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        assert ws["A4"].value == "Restoran / Şube Adı: Restoran Test"
        assert ws["A5"].value == "Açık Adres: Adres Test"
        assert ws["D4"].value == "Denetim Tarihi: 2026-01-15"
        assert ws["D5"].value == "İSG Uzmanı / Denetçi: Test Denetçi"

    def test_11_headers_in_order(self):
        from server import export_excel

        audit = _audit()
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        expected = [
            "No", "Kategori", "Tehlike & Risk Maddesi", "Cevap",
            "Olasılık (O)", "Şiddet (Ş)", "Risk Skoru (R)",
            "Risk Seviyesi", "Sorumlu", "Termin Süresi",
            "Alınması Gereken Tedbir (DÖF)",
        ]
        actual = [ws.cell(row=7, column=i).value for i in range(1, 12)]
        assert actual == expected

    def test_header_fill_and_font(self):
        from openpyxl.styles import PatternFill
        from server import export_excel

        audit = _audit()
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        # Header hücreleri: ilk 11 kolon, 7. satır
        for col in range(1, 12):
            cell = ws.cell(row=7, column=col)
            assert cell.font.bold is True
            assert cell.font.color is not None
            rgb = getattr(cell.font.color, "rgb", None) or getattr(cell.font.color, "value", None)
            assert str(rgb).upper().endswith("FFFFFF")
            assert cell.fill.fill_type == "solid"
            fg = getattr(cell.fill.fgColor, "rgb", None) or getattr(cell.fill.fgColor, "value", None)
            assert str(fg).upper().endswith("8B0000")

    def test_column_widths_strict(self):
        from server import export_excel

        audit = _audit()
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        widths = [6, 24, 55, 12, 12, 12, 14, 18, 20, 15, 50]
        for i, expected_w in enumerate(widths, start=1):
            col_letter = ws.cell(row=1, column=i).column_letter
            actual = ws.column_dimensions[col_letter].width
            assert actual == expected_w, (
                f"Kolon {col_letter} width beklenen {expected_w}, gelen {actual}"
            )


# ===========================================================================
# 2) Veri sözleşmesi — snapshot, override, cevap mapping
# ===========================================================================


class TestDataContract:
    """Veri bütünlüğü: snapshot, override, cevap mapping."""

    def test_turkish_unicode_preserved_no_ascii_transliteration(self):
        from server import export_excel

        audit = _audit()
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        # q1'in "Tehlike & Risk Maddesi" kolonundaki (C) Türkçe metni
        turkish_question = ws.cell(row=8, column=3).value
        # fix_tr_text sonrası "Sisli subesi" / "topraklama olcumu" gibi
        # ASCII-safe hale gelirdi. Burada Unicode karakterlerin durduğunu
        # doğruluyoruz: Ş, ğ, ş, İ, ı, ç, ö, ü, Ö, Ü, Ğ.
        assert "Şişli" in turkish_question, "Ş korunmadı"
        assert "şubesinde" in turkish_question, "ş korunmadı"
        # Sorumlu kolonu (I, kolon 9)
        assert ws.cell(row=8, column=9).value == "İSG Uzmanı — Çağatay Gümüş"

    def test_snapshot_used_not_global_questions(self):
        """Snapshot'tan gelen sorular kullanılır; global QUESTIONS YOK."""
        from server import export_excel

        # Sadece snapshot'taki 5 soru işlenir (84'lük QUESTIONS setinden
        # yalnız bu 5'i basılır).
        audit = _audit()
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        # Header satır 7; ilk veri satırı 8; 5 soru → 8, 9, 10, 11, 12
        for offset, expected_no in enumerate([1, 2, 3, 4, 5], start=0):
            no_cell = ws.cell(row=8 + offset, column=1).value
            assert no_cell == expected_no
        # 6. satır → 6. soru olmamalı (boş olabilir)
        assert ws.cell(row=13, column=1).value in (None, "")

    def test_answer_mapping_evet_hayir_na_empty(self):
        from server import export_excel

        audit = _audit(answers={
            "1": "HAYIR",  # show_risk=True → score 25 / Kabul Edilemez
            "2": "HAYIR",  # override ile değişir — ayrı testte
            "3": "EVET",   # → "Uygun"
            "4": "NA",     # → "N/A"
            "5": "",       # → "Yanıtlanmadı"
        })
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        # Cevap kolonu (kolon 4) — satır 8..12
        assert ws.cell(row=8, column=4).value == "Uygun Değil"      # q1 HAYIR
        assert ws.cell(row=9, column=4).value == "Uygun Değil"      # q2 HAYIR
        assert ws.cell(row=10, column=4).value == "Uygun"           # q3 EVET
        assert ws.cell(row=11, column=4).value == "N/A"             # q4 NA
        assert ws.cell(row=12, column=4).value == "Yanıtlanmadı"    # q5 empty

    def test_override_reflected_in_excel_row(self):
        """Override (5,5) q2 default (3,4) üzerine yazılır; skor 25, level
        kanonik helper ile 'Kabul Edilemez' olur."""
        from server import export_excel

        audit = _audit(
            answers={"1": "HAYIR", "2": "HAYIR"},
            overrides={"2": {"probability": 5, "severity": 5}},
        )
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        # q2 satır 9 — Olasılık(5), Şiddet(5), Risk Skoru(25), Seviye("Kabul Edilemez")
        assert ws.cell(row=9, column=5).value == 5
        assert ws.cell(row=9, column=6).value == 5
        assert ws.cell(row=9, column=7).value == 25
        assert ws.cell(row=9, column=8).value == "Kabul Edilemez"

    def test_risk_color_matches_canonical_helper(self):
        """Risk rengi kanonik ``risk_level`` sabitlerine göre; skor >=15/8
        eşik mantığı YOK."""
        from openpyxl.styles import PatternFill
        from server import export_excel

        audit = _audit(
            answers={"1": "HAYIR", "2": "HAYIR"},
            overrides={"2": {"probability": 5, "severity": 5}},
        )
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        # q1: Kabul Edilemez (skor 25) → DC2626
        cell1 = ws.cell(row=8, column=8)
        assert cell1.fill.fill_type == "solid"
        fg1 = getattr(cell1.fill.fgColor, "rgb", None) or getattr(cell1.fill.fgColor, "value", None)
        assert str(fg1).upper().endswith("DC2626")
        # q2: override ile (5,5) → Kabul Edilemez → DC2626
        cell2 = ws.cell(row=9, column=8)
        fg2 = getattr(cell2.fill.fgColor, "rgb", None) or getattr(cell2.fill.fgColor, "value", None)
        assert str(fg2).upper().endswith("DC2626")

    def test_risk_color_dikkate_deger_yellow(self):
        """Default (3,4) → skor 12 → Dikkate Değer → F59E0B."""
        from server import export_excel

        audit = _audit(answers={"2": "HAYIR"})
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        cell = ws.cell(row=9, column=8)  # q2 satırı
        assert cell.value == "Dikkate Değer"
        fg = getattr(cell.fill.fgColor, "rgb", None) or getattr(cell.fill.fgColor, "value", None)
        assert str(fg).upper().endswith("F59E0B")

    def test_termin_tedbir_only_for_hayir_rows(self):
        """Termin (kolon 10) ve Tedbir (kolon 11) yalnız HAYIR satırlarında
        gerçek değer; aksi halde '-'."""
        from server import export_excel

        audit = _audit(answers={
            "1": "HAYIR", "2": "HAYIR",
            "3": "EVET", "4": "NA", "5": "",
        })
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        # q1 HAYIR → termin "2026-09-30", tedbir içerikli metin
        assert ws.cell(row=8, column=10).value == "2026-09-30"
        assert "Etiketleri güncelle" in (ws.cell(row=8, column=11).value or "")
        # q3 EVET → termin/tedbir "-"
        assert ws.cell(row=10, column=10).value == "-"
        assert ws.cell(row=10, column=11).value == "-"
        # q5 empty → "-"
        assert ws.cell(row=12, column=10).value == "-"
        assert ws.cell(row=12, column=11).value == "-"

    def test_legacy_snapshot_fallback_field_mapping(self):
        """Snapshot'ta yeni şema alanları yoksa (soru/sorumlu/termin/tedbir/
        kategori) legacy fallback kullanılır; export düşürmez."""
        from server import export_excel

        audit = _audit(legacy=True, answers={"1": "HAYIR"})
        with fake_db([_tpl_legacy()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        # Kolon 2 → Kategori (legacy: "Elektrik")
        assert ws.cell(row=8, column=2).value == "Elektrik"
        # Kolon 3 → Soru (legacy: "Pano topraklaması ölçüldü mü?")
        assert ws.cell(row=8, column=3).value == "Pano topraklaması ölçüldü mü?"
        # Kolon 9 → Sorumlu
        assert ws.cell(row=8, column=9).value == "İSG Uzmanı"
        # Kolon 10 → Termin
        assert ws.cell(row=8, column=10).value == "2026-11-30"
        # Kolon 11 → Tedbir
        assert ws.cell(row=8, column=11).value == "Topraklama ölçümü planla."

    def test_no_resmi_onayli_in_status(self):
        """Sözleşme gereği 'RESMİ ONAYLI' string'i status metninde YOKTUR."""
        from server import export_excel

        audit = _audit(is_completed=True)
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
            wb = _load_workbook(resp)
        ws = wb.active
        # A2 — sadece TAMAMLANDI var, RESMİ ONAYLI yok
        assert "RESMİ ONAYLI" not in (ws["A2"].value or "")


# ===========================================================================
# 3) HTTP sözleşmesi — MIME, Content-Disposition, PDF davranışı
# ===========================================================================


class TestHTTPContract:
    """MIME type, Content-Disposition UTF-8, PDF davranışı."""

    def test_mime_type_is_xlsx(self):
        from server import export_excel

        audit = _audit()
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
        assert resp.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def test_content_disposition_utf8_with_isg_prefix(self):
        from server import export_excel

        audit = _audit()
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_excel(audit_id=AID, current_user=_user()))
        cd = resp.headers.get("content-disposition") or ""
        # Hem ASCII hem UTF-8 parça var
        assert "ISG_Risk_Analizi_Restoran_Test.xlsx" in cd
        assert "filename*=UTF-8''" in cd
        # Restaurant name içindeki ASCII fallback (boşluk → _) korunur
        assert "filename=\"ISG_Risk_Analizi_Restoran_Test.xlsx\"" in cd

    def test_pdf_export_unchanged(self):
        """Bu PR'da export_pdf'e dokunulmadı; minimal smoke test."""
        from server import export_pdf

        audit = _audit(answers={"1": "HAYIR"})
        with fake_db([_tpl_tr()]):
            push_audit(__import__("server").db, audit)
            resp = run_async(export_pdf(audit_id=AID, current_user=_user()))
        assert resp.media_type == "application/pdf"
        body = _body_bytes(resp)
        assert body[:4] == b"%PDF"
        # Content-Disposition hala PDF'inki — export_pdf dokunulmadı
        cd = resp.headers.get("content-disposition") or ""
        assert "filename*=UTF-8''" in cd


# ===========================================================================
# 4) Negatif senaryo — audit bulunamadı
# ===========================================================================


class TestNegativeScenarios:
    """Geçersiz ID / bulunamayan audit."""

    def test_invalid_id_returns_400(self):
        from fastapi import HTTPException
        from server import export_excel

        with fake_db([_tpl_tr()]):
            with pytest.raises(HTTPException) as exc:
                run_async(export_excel(audit_id="not-a-valid-id", current_user=_user()))
        assert exc.value.status_code == 400

    def test_missing_audit_returns_404(self):
        from fastapi import HTTPException
        from server import export_excel

        with fake_db([_tpl_tr()]):
            nonexistent = str(ObjectId("507f1f77bcf86cd799439099"))
            with pytest.raises(HTTPException) as exc:
                run_async(export_excel(audit_id=nonexistent, current_user=_user()))
        assert exc.value.status_code == 404
